import sys
from pathlib import Path
sys.path[:0] = [str(Path(__file__).resolve().parents[1]), str(Path(__file__).resolve().parents[1] / "tools")]
"""Synthetic workbook fixtures; no private source workbook is committed."""
import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from PySide6.QtWidgets import QApplication
from PySide6.QtCore import Qt
from importers import import_agenda, agenda_events
from main import RowsEditor, AgendaImportDialog, RUNDOWN_COLUMNS

class AgendaTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    def preset(self, path):
        book=Workbook();s=book.active;s.title='Rundown'
        s.append(['RUNDOWN KEGIATAN']);s.append(['TRIP OBSERVASI']);s.append(['Rabu, 14 Oktober 2026'])
        s.append(['NO','WAKTU','KEGIATAN','TEMPAT','NARASUMBER','PENANGGUNG JAWAB'])
        s.append([1,'21:30 - 04.00','Istirahat peserta','Rumah','Guru','Mentor'])
        s.append([2,'21:30 - 22:30','Evaluasi panitia','Saung',None,None])
        s.merge_cells('E5:E6');s.merge_cells('F5:F6')
        s.append([3,'22:30','Istirahat panitia','Rumah','-','-'])
        s.append(['RUNDOWN KEGIATAN']);s.append(['TRIP OBSERVASI']);s.append(['Kamis, 15 Oktober 2026'])
        s.append(['NO','WAKTU','KEGIATAN','TEMPAT','NARASUMBER','PENANGGUNG JAWAB'])
        s.append([1,'04:00 - 05:30','Bangun tidur','Rumah','Guru','Mentor'])
        private=book.create_sheet('Panitia');private.append(['NO','NAMA','TELEPON']);private.append([1,'Private person','0812-private']);private['D2']='=1+1'
        book.save(path);return book
    def test_multi_day_merged_credits_parallel_overnight_and_private_sheet(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'preset.xlsx';self.preset(path);blocks=import_agenda(path)
            self.assertEqual(len(blocks),2);first=blocks[0]
            self.assertTrue(first['endNextDay']);self.assertEqual(first['end'],'04:00')
            self.assertEqual(first['rundown'][1]['speaker'],'Guru');self.assertEqual(first['rundown'][1]['coordinator'],'Mentor')
            self.assertEqual(first['rundown'][2]['end'],'');self.assertNotIn('0812-private',str(blocks))
            events=agenda_events(blocks);self.assertEqual(events,agenda_events(blocks));self.assertFalse(events[0]['notify'])
            self.assertIn('Hari 2',agenda_events([blocks[1]])[0]['title'])
    def test_formula_date_mismatch_and_huge_merge_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'preset.xlsx'
            for mode in ('formula','date','merge'):
                book=self.preset(path);s=book.active
                if mode=='formula':s['C5']='=HYPERLINK("https://example.com")'
                elif mode=='date':s['A2']='TRIP OBSERVASI';s['A1']='RUNDOWN KEGIATAN';s.insert_rows(4);s['A4']='15 Oktober 2026'
                else:s.merge_cells('A100:A2001')
                book.save(path)
                with self.assertRaises(ValueError):import_agenda(path)
    def test_simple_format_rejects_extra_columns_and_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'simple.xlsx';b=Workbook();s=b.active;s.append(['Mulai','Selesai','Agenda','Lokasi']);s.append(['07:00','08:00','Agenda','Aula']);s['H2']='Unmapped private information';b.save(path)
            with self.assertRaises(ValueError):import_agenda(path)
    def test_preview_edits_preserve_checkbox_and_recalculate_end(self):
        rows=[dict(start='21:30',end='04:00',endNextDay=True,agenda='Rest',location='Home',speaker='Teacher',coordinator='Mentor'),dict(start='22:30',end='',agenda='Marker',location='Home')]
        editor=RowsEditor(RUNDOWN_COLUMNS,rows);self.assertEqual(editor.values()[0]['endNextDay'],True);self.assertEqual(editor.values()[1]['endNextDay'],False)
        editor.add_row();last=editor.table.rowCount()-1
        self.assertEqual(editor.table.item(last,2).checkState(),Qt.CheckState.Unchecked)
        dialog=AgendaImportDialog([dict(title='Trip',date='2026-10-14',rundown=rows,warnings=[])],None,multi=True)
        self.assertTrue(dialog.selected_blocks()[0]['endNextDay']);dialog.reject();editor.close()

if __name__=='__main__':unittest.main()
