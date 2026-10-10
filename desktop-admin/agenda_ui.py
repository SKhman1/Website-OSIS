"""Native selection, manual mapping and optional authenticated AI assistance."""
import copy
import json
from PySide6.QtCore import QSettings, QTimer
from PySide6.QtWidgets import (QHBoxLayout,QVBoxLayout,QFormLayout,QGridLayout,QLineEdit,QComboBox,QSpinBox,QCheckBox,QPushButton,QTableWidget,QTableWidgetItem,QHeaderView,QScrollArea,QWidget)
from security_ui import ApiDialog, caption
from agenda_mapping import sheets,read_grid,bounds,ai_payload,detect_plans,mapped_blocks,ROLES,redact
from openpyxl.utils.cell import get_column_letter

CAPTIONS={"period":"Waktu gabungan","start":"Jam mulai","end":"Jam selesai","agenda":"Agenda/kegiatan","location":"Tempat","speaker":"Narasumber","coordinator":"Penanggung jawab"}

class FlexibleAgendaDialog(ApiDialog):
    def __init__(self,parent,client,path):
        super().__init__(parent,client,"Pemetaan agenda Excel");self.resize(1020,760);self.path=path;self.grid=None;self.plan={"blocks":[]};self.active=-1;self.output=[];self.loading=False
        outer=self.layout;panel=QWidget();self.layout=QVBoxLayout(panel);scroll=QScrollArea();scroll.setWidgetResizable(True);scroll.setWidget(panel);outer.addWidget(scroll)
        self.settings=QSettings("VimsatiNattaDarsa","AgendaPresets")
        self.layout.addWidget(caption("Pilih hanya sheet dan area rundown. Data di luar area tidak dikirim. Nomor telepon/email yang dikenali disamarkan pada preview AI. Formula tidak dijalankan."))
        top=QHBoxLayout();self.sheet=QComboBox();self.sheet.addItems(sheets(path));self.area=QLineEdit("A1:H60");self.area.setPlaceholderText("A1:H60 · maksimal 200 baris, 20 kolom")
        top.addWidget(self.sheet);top.addWidget(self.area);load=QPushButton("Baca area lokal");load.clicked.connect(self.load_area);top.addWidget(load);self.layout.addLayout(top)
        self.table=QTableWidget();self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers);self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive);self.table.setMinimumHeight(160);self.layout.addWidget(self.table,1)
        self.consent=QCheckBox("Saya sudah memeriksa preview dan mengizinkan pengiriman area ini ke Groq.")
        self.layout.addWidget(self.consent)
        actions=QHBoxLayout()
        for text,callback in [("Deteksi kolom lokal",self.detect),("Bantu baca dengan AI",self.infer),("Tambah tabel manual",self.add_manual)]:
            item=QPushButton(text);item.clicked.connect(callback);actions.addWidget(item)
        self.layout.addLayout(actions)
        self.pick=QComboBox();self.pick.currentIndexChanged.connect(self.choose_block);self.layout.addWidget(self.pick)
        form=QFormLayout();self.first=QSpinBox();self.last=QSpinBox();self.first.setRange(1,2000);self.last.setRange(1,2000);self.title_ref=QLineEdit();self.date_ref=QLineEdit()
        self.title_ref.setPlaceholderText("Sel judul, contoh A3; kosong = isi di preview berikutnya")
        self.date_ref.setPlaceholderText("Sel tanggal lengkap; kosong = isi di preview berikutnya")
        rows=QHBoxLayout();rows.addWidget(caption("Baris mulai"));rows.addWidget(self.first);rows.addWidget(caption("Baris akhir"));rows.addWidget(self.last);form.addRow("Baris isi (tanpa header)",rows);form.addRow("Sel judul",self.title_ref);form.addRow("Sel tanggal",self.date_ref)
        cols=QGridLayout();self.columns={}
        for i,role in enumerate(ROLES):
            box=QComboBox();self.columns[role]=box;cols.addWidget(caption(CAPTIONS[role]),i//2,(i%2)*2);cols.addWidget(box,i//2,(i%2)*2+1)
        form.addRow(cols);self.layout.addLayout(form)
        preset=QHBoxLayout();self.preset_name=QLineEdit();self.preset_name.setPlaceholderText("Nama preset (hanya pemetaan, tanpa isi Excel)")
        self.presets=QComboBox();self.presets.addItems(self.settings.childGroups());preset.addWidget(self.preset_name);save=QPushButton("Simpan preset");save.clicked.connect(self.save_preset);preset.addWidget(save);preset.addWidget(self.presets);restore=QPushButton("Pakai preset");restore.clicked.connect(self.restore_preset);preset.addWidget(restore);self.layout.addLayout(preset)
        finish=QHBoxLayout();cancel=QPushButton("Batal");cancel.clicked.connect(self.reject);finish.addWidget(cancel);finish.addStretch();apply=QPushButton("Periksa hasil pemetaan");apply.clicked.connect(self.finish);finish.addWidget(apply);self.layout.addLayout(finish)
        self.load_area()
    def load_area(self):
        try:
            self.grid=read_grid(self.path,self.sheet.currentText(),self.area.text().strip().upper());left,top,right,bottom=bounds(self.grid["range"])
            self.plan={"blocks":[]};self.active=-1;self.pick.clear();self.consent.setChecked(False)
            self.table.setRowCount(bottom-top+1);self.table.setColumnCount(right-left+1);self.table.clearContents();self.table.setHorizontalHeaderLabels([get_column_letter(c) for c in range(left,right+1)]);self.table.setVerticalHeaderLabels([str(r) for r in range(top,bottom+1)])
            for cell in self.grid["cells"]:
                c,r,_,_=bounds(cell["ref"]+":"+cell["ref"]);self.table.setItem(r-top,c-left,QTableWidgetItem(redact(cell["value"])))
            for box in self.columns.values():
                box.clear();box.addItem("Tidak ada",0)
                for c in range(left,right+1):box.addItem(get_column_letter(c),c)
            self.status.setText("Area dibaca lokal. Preview menampilkan teks yang akan dikirim jika bantuan AI dipilih.")
            try:self.set_plan(detect_plans(self.grid))
            except ValueError:pass
        except Exception as exc:self.grid=None;self.status.setText(str(exc))
    def stash(self):
        if 0<=self.active<len(self.plan["blocks"]):self.plan["blocks"][self.active]=dict(firstRow=self.first.value(),lastRow=self.last.value(),titleCell=self.title_ref.text().strip().upper(),dateCell=self.date_ref.text().strip().upper(),columns={role:box.currentData() for role,box in self.columns.items()})
    def choose_block(self,index):
        if self.loading:return
        self.stash();self.active=index
        if index<0:return
        block=self.plan["blocks"][index];self.first.setValue(block["firstRow"]);self.last.setValue(block["lastRow"]);self.title_ref.setText(block["titleCell"]);self.date_ref.setText(block["dateCell"])
        for role,box in self.columns.items():box.setCurrentIndex(max(0,box.findData(block["columns"][role])))
    def set_plan(self,plan):
        self.loading=True;self.plan=copy.deepcopy(plan);self.active=-1;self.pick.clear();self.pick.addItems([f"Tabel {i+1} · baris {b['firstRow']}–{b['lastRow']}" for i,b in enumerate(self.plan["blocks"])]);self.loading=False;self.choose_block(0)
    def detect(self):
        try:
            if not self.grid:raise ValueError("Baca area lokal terlebih dahulu.")
            self.set_plan(detect_plans(self.grid));self.status.setText("Kolom terdeteksi lokal. Periksa batas tabel, judul, tanggal, dan kolom.")
        except ValueError as exc:self.status.setText(str(exc))
    def add_manual(self):
        if not self.grid:return
        self.stash();_,top,_,bottom=bounds(self.grid["range"]);self.plan["blocks"].append(dict(firstRow=top,lastRow=bottom,titleCell="",dateCell="",columns={r:0 for r in ROLES}));self.set_plan(self.plan);self.pick.setCurrentIndex(len(self.plan["blocks"])-1)
    def infer(self):
        try:
            if not self.grid:raise ValueError("Baca area lokal terlebih dahulu.")
            if not self.consent.isChecked():raise ValueError("Periksa area yang akan dikirim, lalu centang persetujuan bantuan AI.")
            payload=ai_payload(self.grid)
            def done(result):self.set_plan(result["plan"]);self.status.setText("Pemetaan dari GPT-OSS 20B siap diperiksa. AI tidak mengubah nilai sel atau memublikasikan jadwal.")
            self.call("/api/manage/agenda-ai",payload,done)
        except ValueError as exc:self.status.setText(str(exc))
    def save_preset(self):
        try:
            self.stash();name=self.preset_name.text().strip()
            if not name or len(name)>50 or any(c in name for c in "/\\"):raise ValueError("Isi nama preset maksimal 50 karakter tanpa garis miring.")
            if not self.plan["blocks"]:raise ValueError("Isi pemetaan terlebih dahulu.")
            self.settings.setValue(name+"/plan",json.dumps(self.plan));self.presets.clear();self.presets.addItems(self.settings.childGroups());self.status.setText("Preset disimpan lokal; isi sel Excel dan API key tidak disimpan.")
        except ValueError as exc:self.status.setText(str(exc))
    def restore_preset(self):
        try:
            plan=json.loads(self.settings.value(self.presets.currentText()+"/plan","{}"));self.set_plan(plan);self.status.setText("Preset dimuat. Periksa kolom dan batas baris karena file dapat berubah.")
        except (ValueError,KeyError,TypeError):self.status.setText("Preset belum tersedia atau tidak valid.")
    def finish(self):
        try:
            if not self.grid:raise ValueError("Baca area terlebih dahulu.")
            self.stash();self.output=mapped_blocks(self.grid,self.plan);self.accept()
        except ValueError as exc:self.status.setText(str(exc))

def flexible_blocks(parent,client,path):
    dialog=FlexibleAgendaDialog(parent,client,path)
    return dialog.output if dialog.exec()==dialog.DialogCode.Accepted else None
