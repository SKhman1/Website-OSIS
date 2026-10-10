"""Flexible, local Excel mapping. AI returns coordinates, never replacement text."""
import json
import re
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter
from importers import xlsx_metadata, clock, rundown_bounds, MONTHS

ROLES=("period","start","end","agenda","location","speaker","coordinator")
ALIASES={"period":("waktu","jam","time","jam kegiatan"),"start":("mulai","jam mulai","waktu mulai","start","start time"),"end":("selesai","jam selesai","waktu selesai","end","end time"),"agenda":("agenda","kegiatan","acara","activity","aktivitas"),"location":("lokasi","tempat","location","venue"),"speaker":("narasumber","pemateri","speaker"),"coordinator":("penanggung jawab","pj","pic","koordinator","coordinator")}

def redact(value):
    value=re.sub(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}","[email disamarkan]",value,flags=re.I)
    return re.sub(r"(?:\+?62|0)\d[\d ()-]{6,}\d",lambda m:"[nomor disamarkan]" if 10<=len(re.sub(r"\D","",m[0]))<=15 else m[0],value)

def sheets(path):
    if Path(path).suffix.lower()!=".xlsx":raise ValueError("Pemetaan fleksibel memakai .xlsx.")
    xlsx_metadata(path)
    book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        if len(book.worksheets)>20:raise ValueError("Maksimal 20 sheet.")
        return [s.title for s in book if s.sheet_state=="visible"]
    finally:book.close()

def bounds(area):
    if not re.fullmatch(r"[A-Z]{1,2}[1-9]\d{0,3}:[A-Z]{1,2}[1-9]\d{0,3}",area):raise ValueError("Isi area seperti A1:H60.")
    left,top,right,bottom=range_boundaries(area)
    if right<left or bottom<top or right>64 or bottom>2000 or right-left>=20 or bottom-top>=200:raise ValueError("Area maksimal 200 baris dan 20 kolom, di dalam 2000 baris/64 kolom.")
    return left,top,right,bottom

def cell_text(value):
    if value is None:return ""
    if isinstance(value,time):return value.strftime("%H:%M")
    if isinstance(value,(datetime,date)):return value.strftime("%Y-%m-%d")
    if isinstance(value,float) and 0<=value<1:
        minutes=round(value*1440)
        if minutes<1440:return f"{minutes//60:02d}:{minutes%60:02d}"
    return str(value).strip()

def read_grid(path,sheet,area):
    left,top,right,bottom=bounds(area);merges=xlsx_metadata(path)
    book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        if sheet not in book.sheetnames or book[sheet].sheet_state!="visible":raise ValueError("Sheet tidak tersedia.")
        result={}
        for row in book[sheet].iter_rows(min_row=top,max_row=bottom,min_col=left,max_col=right):
            for cell in row:
                if cell.data_type=="f":raise ValueError("Area terpilih berisi formula. Gunakan salinan nilai biasa.")
                text=cell_text(cell.value)
                if text:result[cell.coordinate]=text
        # Only explicit merged areas inherit values; never inherit ordinary blanks.
        for a,b,c,d in merges.get(sheet,[]):
            if left<=a<=right and top<=b<=bottom:
                value=result.get(f"{get_column_letter(a)}{b}")
                if value:
                    for r in range(b,min(d,bottom)+1):
                        for col in range(a,min(c,right)+1):result[f"{get_column_letter(col)}{r}"]=value
        return {"sheet":sheet,"range":area,"cells":[{"ref":key,"value":value} for key,value in result.items()]}
    finally:book.close()

def ai_payload(grid):
    output={**grid,"sheet":redact(grid["sheet"]),"cells":[{"ref":c["ref"],"value":redact(c["value"])} for c in grid["cells"]]}
    if not output["cells"] or any(len(c["value"])>500 for c in output["cells"]) or len(json.dumps(output,ensure_ascii=False,separators=(",",":")))>24000:raise ValueError("Pilih area lebih kecil untuk AI (maksimal 24.000 karakter; 500 karakter per sel).")
    return output

def parse_date(value):
    found=re.search(r"\b(\d{4})-(\d{2})-(\d{2})\b",value)
    if found:return date(*map(int,found.groups())).isoformat()
    found=re.search(r"\b(\d{1,2})[/.](\d{1,2})[/.](\d{4})\b",value)
    if found:return date(int(found[3]),int(found[2]),int(found[1])).isoformat()
    found=re.search(r"\b(\d{1,2})\s+("+"|".join(MONTHS)+r")\s+(\d{4})\b",value,re.I)
    if found:return date(int(found[3]),MONTHS[found[2].lower()],int(found[1])).isoformat()
    return ""

def detect_plans(grid):
    values={c["ref"]:c["value"] for c in grid["cells"]};left,top,right,bottom=bounds(grid["range"]);headers=[]
    for row in range(top,bottom+1):
        columns={role:0 for role in ROLES}
        for col in range(left,right+1):
            text=re.sub(r"[^a-z0-9 ]","",values.get(f"{get_column_letter(col)}{row}","").lower()).strip()
            for role,aliases in ALIASES.items():
                if text in aliases:columns[role]=col
        if columns["agenda"] and (columns["period"] or columns["start"]):
            if columns["start"]:columns["period"]=0
            elif columns["period"]:columns["end"]=0
            headers.append((row,columns))
    blocks=[]
    for index,(row,columns) in enumerate(headers):
        last=headers[index+1][0]-1 if index+1<len(headers) else bottom
        while last>row and not any(values.get(f"{get_column_letter(col)}{last}") for col in columns.values() if col):last-=1
        context=[(key,value) for key,value in values.items() if max(top,row-10)<=range_boundaries(key)[1]<row]
        dates=[key for key,value in context if parse_date(value)]
        titles=[key for key,value in context if not parse_date(value) and len(value)<=160 and value.upper() not in ("RUNDOWN KEGIATAN","RUNDOWN","HARI","TANGGAL")]
        if last>row:blocks.append(dict(firstRow=row+1,lastRow=last,titleCell=titles[0] if titles else "",dateCell=dates[-1] if dates else "",columns=columns))
    if not blocks:raise ValueError("Kolom belum dikenali. Pilih kolom manual atau gunakan bantuan AI.")
    return {"blocks":blocks}

def mapped_blocks(grid,plan):
    left,top,right,bottom=bounds(grid["range"]);values={c["ref"]:c["value"] for c in grid["cells"]};output=[];previous=top-1
    if set(plan)!={"blocks"} or not isinstance(plan["blocks"],list) or not 1<=len(plan["blocks"])<=12:raise ValueError("Pemetaan tidak valid.")
    for block in plan["blocks"]:
        if set(block)!={"firstRow","lastRow","titleCell","dateCell","columns"}:raise ValueError("Kolom pemetaan tidak valid.")
        first,last=block["firstRow"],block["lastRow"];cols=block["columns"]
        if type(first) is not int or type(last) is not int or not previous<first<=last<=bottom or set(cols)!=set(ROLES):raise ValueError("Baris pemetaan tidak valid.")
        previous=last
        if any(type(c) is not int or (c!=0 and not left<=c<=right) for c in cols.values()) or not cols["agenda"] or bool(cols["period"])==bool(cols["start"]) or (cols["period"] and cols["end"]):raise ValueError("Pilih waktu gabungan atau kolom mulai/selesai dan kolom agenda.")
        used=[c for c in cols.values() if c]
        if len(used)!=len(set(used)):raise ValueError("Satu kolom hanya untuk satu fungsi.")
        if any(not isinstance(block[k],str) or (block[k] and block[k] not in values) for k in ("titleCell","dateCell")):raise ValueError("Sel judul/tanggal tidak ditemukan.")
        rows=[];sources=[]
        for r in range(first,last+1):
            get=lambda role:values.get(f"{get_column_letter(cols[role])}{r}","") if cols[role] else ""
            if not get("agenda") and not (get("period") or get("start")):continue
            period=get("period");parts=re.split(r"\s*[-–—]\s*",period) if period else [get("start"),get("end")]
            if not 1<=len(parts)<=2:raise ValueError(f"Baris {r}: waktu tidak valid.")
            def normalize(value):
                value=re.sub(r"\s*WIB\s*","",value,flags=re.I).strip().replace(".",":")
                if re.fullmatch(r"\d:[0-5]\d",value):value="0"+value
                return clock(value)
            try:a=normalize(parts[0]);b=normalize(parts[1]) if len(parts)>1 and parts[1] else ""
            except ValueError:raise ValueError(f"Baris {r}: periksa jam asli atau batas baris tabel.") from None
            rows.append(dict(start=a,end=b,endNextDay=bool(b and b<a),agenda=get("agenda"),location=get("location"),speaker=get("speaker"),coordinator=get("coordinator")));sources.append(f"{grid['sheet']} · baris {r}")
        if len(rows)>200:raise ValueError("Maksimal 200 sesi per hari.")
        try:period=rundown_bounds(rows)
        except ValueError as exc:raise ValueError(f"Periksa tabel baris {first}–{last}: {exc}") from None
        output.append(dict(title=values.get(block["titleCell"],""),date=parse_date(values.get(block["dateCell"],"")),rundown=rows,warnings=["Pemetaan perlu diperiksa sebelum publikasi. Isi berasal dari sel Excel asli."],source=grid["sheet"],sourceRows=sources,**period))
    return output
